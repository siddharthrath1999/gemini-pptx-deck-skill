#!/usr/bin/env python3
"""Read and edit an existing .pptx while keeping its own theme, masters and formatting.

Usage:
  python3 edit_deck.py inspect deck.pptx [-o outline.json]
      Prints a readable outline and writes every slide, shape, text, table, chart and the deck's
      theme to outline.json. Use the slide "id" values from it in edits; they never change.
  python3 edit_deck.py apply deck.pptx edits.json -o edited.pptx
      Applies a list of edit operations (see references/editing.md) and checks the result.
Exit code 0 = applied cleanly, 1 = applied with problems to fix, 2 = could not apply.
"""

import argparse
import copy
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402
from pptx.util import Emu  # noqa: E402

import dk_lint  # noqa: E402
import dk_measure as M  # noqa: E402
import dk_units  # noqa: E402

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
DESIGN_W = 13.333


def inches(v):
    return round(Emu(v).inches, 2) if v is not None else None


# ---- reading ------------------------------------------------------------------------------------
def deck_theme(prs):
    out = {"colors": {}, "fonts": {}}
    try:
        part = prs.slide_masters[0].part.part_related_by(RT.THEME)
        root = etree.fromstring(part.blob)
        scheme = root.find(".//a:clrScheme", NS)
        for el in (scheme if scheme is not None else []):
            name = etree.QName(el).localname
            srgb = el.find("a:srgbClr", NS)
            sysc = el.find("a:sysClr", NS)
            val = srgb.get("val") if srgb is not None else (sysc.get("lastClr") if sysc is not None else None)
            if val:
                out["colors"][name] = val.upper()
        for key, tag in (("major", "a:majorFont"), ("minor", "a:minorFont")):
            latin = root.find(".//%s/a:latin" % tag, NS)
            if latin is not None:
                out["fonts"][key] = latin.get("typeface")
    except Exception as exc:
        out["error"] = str(exc)
    return out


def shape_info(shape):
    info = {"id": shape.shape_id, "name": shape.name, "x": inches(shape.left), "y": inches(shape.top),
            "w": inches(shape.width), "h": inches(shape.height)}
    if shape.is_placeholder:
        info["placeholder"] = str(shape.placeholder_format.type).split(".")[-1].split(" ")[0].lower()
    if getattr(shape, "has_chart", False) and shape.has_chart:
        ch = shape.chart
        data = {"type": str(ch.chart_type).split(".")[-1].split(" ")[0].lower(), "series": []}
        try:
            data["categories"] = [str(c) for c in ch.plots[0].categories]
        except Exception:
            data["categories"] = []
        for plot in ch.plots:
            for ser in plot.series:
                data["series"].append({"name": ser.name, "values": list(ser.values)})
        info["kind"], info["chart"] = "chart", data
    elif getattr(shape, "has_table", False) and shape.has_table:
        info["kind"] = "table"
        info["table"] = [[cell.text for cell in row.cells] for row in shape.table.rows]
    elif shape.shape_type is not None and "PICTURE" in str(shape.shape_type):
        info["kind"] = "picture"
    elif shape.shape_type is not None and "GROUP" in str(shape.shape_type):
        info["kind"] = "group"
        info["children"] = [shape_info(s) for s in shape.shapes]
    elif shape.has_text_frame:
        info["kind"] = "text"
        info["text"] = [p.text for p in shape.text_frame.paragraphs]
        sizes = [seg.font.size.pt for p in shape.text_frame.paragraphs for seg in p.runs if seg.font.size]
        if sizes:
            info["font_size"] = max(sizes)
    else:
        info["kind"] = "shape"
    return info


def slide_title(slide):
    if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
        return slide.shapes.title.text_frame.text
    best, size = "", 0
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            s = max([seg.font.size.pt for p in sh.text_frame.paragraphs for seg in p.runs if seg.font.size] + [0])
            if s > size:
                best, size = sh.text_frame.text, s
    return best


def inspect(prs, path):
    w, h = inches(prs.slide_width), inches(prs.slide_height)
    out = {"file": os.path.basename(path), "slide_size_in": [w, h],
           "aspect": "16:9" if abs(w / h - 16 / 9) < 0.02 else ("4:3" if abs(w / h - 4 / 3) < 0.02 else "other"),
           "theme": deck_theme(prs), "slides": []}
    for n, slide in enumerate(prs.slides, 1):
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        out["slides"].append({"n": n, "id": slide.slide_id, "layout": slide.slide_layout.name, "title": slide_title(slide),
                              "notes": notes, "shapes": [shape_info(s) for s in slide.shapes]})
    return out


def print_outline(info):
    print("%s · %d slides · %s × %s in (%s)" % (info["file"], len(info["slides"]), info["slide_size_in"][0], info["slide_size_in"][1], info["aspect"]))
    t = info["theme"]
    print("Theme fonts: %s / %s · colours: %s" % (t["fonts"].get("major"), t["fonts"].get("minor"),
                                                   ", ".join("%s %s" % kv for kv in t["colors"].items())))
    for s in info["slides"]:
        kinds = {}
        for sh in s["shapes"]:
            kinds[sh["kind"]] = kinds.get(sh["kind"], 0) + 1
        print("  %2d · id %d · %s · \"%s\" · %s%s" % (s["n"], s["id"], s["layout"], (s["title"] or "").replace("\n", " ")[:70],
                                                   ", ".join("%d %s" % (v, k) for k, v in kinds.items()), " · notes" if s["notes"] else ""))


# ---- editing helpers ----------------------------------------------------------------------------
class EditError(Exception):
    pass


def slides_by_id(prs):
    return {s.slide_id: s for s in prs.slides}


def resolve_slide(prs, ref, original):
    """ref is a slide id from inspect, or {"n": position in the original deck}."""
    if isinstance(ref, dict) and "n" in ref:
        ref = original.get(int(ref["n"]))
    by_id = slides_by_id(prs)
    if ref not in by_id:
        raise EditError("No slide with id %r. Use the ids from edit_deck.py inspect." % (ref,))
    return by_id[ref]


def resolve_shape(slide, ref):
    for sh in slide.shapes:
        if sh.shape_id == ref or sh.name == ref:
            return sh
    raise EditError("No shape %r on slide id %d. Shapes: %s" % (ref, slide.slide_id, ", ".join("%d %s" % (s.shape_id, s.name) for s in slide.shapes)))


def write_paragraphs(text_frame, paragraphs):
    """Replace text but keep the first run's look and the first paragraph's spacing and bullets."""
    paras = list(text_frame.paragraphs)
    first = paras[0]
    p_pr = first._p.find("{%s}pPr" % NS["a"])
    first_seg = first.runs[0] if first.runs else None
    r_pr = first_seg._r.find("{%s}rPr" % NS["a"]) if first_seg is not None else None
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    for child in list(first._p):
        if etree.QName(child).localname in ("r", "br", "fld"):
            first._p.remove(child)
    for i, value in enumerate(paragraphs):
        para = first if i == 0 else text_frame.add_paragraph()
        if i > 0 and p_pr is not None:
            para._p.insert(0, copy.deepcopy(p_pr))
        seg = para.add_run()
        seg.text = str(value)
        if r_pr is not None:
            old = seg._r.find("{%s}rPr" % NS["a"])
            if old is not None:
                seg._r.remove(old)
            seg._r.insert(0, copy.deepcopy(r_pr))


def text_fit_warning(shape, slide_no):
    if not shape.has_text_frame or not shape.width or not shape.height:
        return None
    paras = [p.text for p in shape.text_frame.paragraphs]
    sizes = [seg.font.size.pt for p in shape.text_frame.paragraphs for seg in p.runs if seg.font.size]
    size = max(sizes) if sizes else 18.0
    fam = None
    for p in shape.text_frame.paragraphs:
        for seg in p.runs:
            fam = fam or seg.font.name
    width = Emu(shape.width).inches - 0.2
    need, _ = M.text_height(paras, fam or "Arial", size, max(0.5, width), False, 1.2)
    if need > Emu(shape.height).inches * 1.08 + 0.05:
        return {"slide": slide_no, "level": "warn", "msg": "Edited text in '%s' may overflow its box (needs about %.1f in, box is %.1f in). Shorten it or check the render."
                % (shape.name, need, Emu(shape.height).inches)}
    return None


def move_after(prs, slide, after_id):
    lst = prs.slides._sldIdLst
    entries = list(lst)
    mine = next(e for e in entries if int(e.get("id")) == slide.slide_id)
    lst.remove(mine)
    if not after_id:
        lst.insert(0, mine)
        return
    anchor = next((e for e in list(lst) if int(e.get("id")) == after_id), None)
    if anchor is None:
        raise EditError("No slide with id %r to place after." % (after_id,))
    anchor.addnext(mine)


def delete_slide(prs, slide):
    lst = prs.slides._sldIdLst
    for e in list(lst):
        if int(e.get("id")) == slide.slide_id:
            rid = e.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            lst.remove(e)
            prs.part.drop_rel(rid)
            return


def renumber_parts(prs):
    """Give every slide and notes part a unique name in deck order (deleting and adding slides can repeat names)."""
    from pptx.opc.packuri import PackURI
    for i, slide in enumerate(prs.slides, 1):
        slide.part.partname = PackURI("/ppt/slides/slide%d.xml" % i)
        if slide.has_notes_slide:
            slide.notes_slide.part.partname = PackURI("/ppt/notesSlides/notesSlide%d.xml" % i)


def theme_from_deck(info, overrides):
    c, f = info["theme"]["colors"], info["theme"]["fonts"]
    accents = [c[k] for k in ("accent1", "accent2", "accent3", "accent4") if k in c]
    t = {"bg": c.get("lt1", "FFFFFF"), "ink": c.get("dk1", "141414"), "dark": c.get("dk2", c.get("dk1", "141414")),
         "accents": accents or ["2F6FED"], "series": [c[k] for k in ("accent1", "accent2", "accent3", "accent4", "accent5", "accent6") if k in c] or None,
         "display": f.get("major") or "Arial", "body": f.get("minor") or "Arial"}
    t = {k: v for k, v in t.items() if v}
    t.update(overrides or {})
    return t


def insert_slides(prs, info, op, original, findings):
    from dk_canvas import Deck
    from dk_layouts import LAYOUTS, NO_FOOTER
    if info["aspect"] != "16:9":
        raise EditError("New slides can only be added to 16:9 decks; this deck is %s. Edit text, tables and charts instead." % info["aspect"])
    dk_units.set_scale(Emu(prs.slide_width).inches / DESIGN_W)
    try:
        deck = Deck(theme_from_deck(info, op.get("theme")), {"page_numbers": False, "footer": ""}, prs=prs)
        for w in deck.theme.warnings:
            if "safe list" not in w:
                findings.append({"slide": 0, "level": "info", "msg": "Inserted slides: " + w})
        anchor = op.get("after", 0)
        if isinstance(anchor, dict):
            anchor = original.get(int(anchor["n"]))
        made = []
        for spec in op.get("slides") or []:
            name = spec.get("layout")
            if name not in LAYOUTS:
                raise EditError("Unknown layout %r in insert_slides." % name)
            default_tone = {"section": "dark", "closing": "dark", "question": "accent", "statement": "accent"}.get(name, "light")
            canvas = deck.new_slide(spec.get("tone") or default_tone, spec)
            LAYOUTS[name](canvas, spec)
            if name not in NO_FOOTER:
                canvas.footer(spec.get("source"), page=False)
            canvas.notes(spec.get("notes"))
            move_after(prs, canvas.slide, anchor)
            anchor = canvas.slide.slide_id
            made.append(canvas.slide.slide_id)
            for f in dk_lint.lint_geometry(canvas) + dk_lint.lint_copy(spec, 0):
                f["msg"] = "New slide id %d: %s" % (canvas.slide.slide_id, f["msg"])
                findings.append(f)
        return made
    finally:
        dk_units.set_scale(1.0)


# ---- apply --------------------------------------------------------------------------------------
def apply(prs, ops, findings):
    info = inspect(prs, "deck")
    original = {s["n"]: s["id"] for s in info["slides"]}
    log = []
    for k, op in enumerate(ops, 1):
        kind = op.get("op")
        try:
            if kind == "set_text":
                slide = resolve_slide(prs, op["slide"], original)
                shape = resolve_shape(slide, op["shape"])
                if not shape.has_text_frame:
                    raise EditError("Shape %r has no text." % op["shape"])
                value = op["text"]
                write_paragraphs(shape.text_frame, value if isinstance(value, list) else str(value).split("\n"))
                w = text_fit_warning(shape, k)
                if w:
                    findings.append(w)
                log.append("set text of '%s' on slide id %d" % (shape.name, slide.slide_id))
            elif kind == "replace_text":
                find, new = op["find"], op["with"]
                targets = [resolve_slide(prs, s, original) for s in op.get("slides", [])] or list(prs.slides)
                count = 0
                for slide in targets:
                    frames = []
                    for sh in slide.shapes:
                        if sh.has_text_frame:
                            frames.append(sh.text_frame)
                        if getattr(sh, "has_table", False) and sh.has_table:
                            frames.extend(cell.text_frame for row in sh.table.rows for cell in row.cells)
                    for tf in frames:
                        for p in tf.paragraphs:
                            for seg in p.runs:
                                if find in seg.text:
                                    count += seg.text.count(find)
                                    seg.text = seg.text.replace(find, new)
                if not count:
                    findings.append({"slide": 0, "level": "warn", "msg": "replace_text found no '%s'. Text split across formatting runs is not matched; use set_text." % find})
                log.append("replaced '%s' %d times" % (find, count))
            elif kind == "set_notes":
                slide = resolve_slide(prs, op["slide"], original)
                slide.notes_slide.notes_text_frame.text = str(op["text"])
                log.append("set notes on slide id %d" % slide.slide_id)
            elif kind == "set_table_cell":
                slide = resolve_slide(prs, op["slide"], original)
                shape = resolve_shape(slide, op["shape"])
                cell = shape.table.cell(int(op["row"]), int(op["col"]))
                write_paragraphs(cell.text_frame, [str(op["text"])])
                log.append("set table cell r%s c%s on slide id %d" % (op["row"], op["col"], slide.slide_id))
            elif kind == "set_chart_data":
                slide = resolve_slide(prs, op["slide"], original)
                shape = resolve_shape(slide, op["shape"])
                data = CategoryChartData()
                data.categories = op["categories"]
                for ser in op["series"]:
                    if len(ser["values"]) != len(op["categories"]):
                        raise EditError("Series %r has %d values for %d categories." % (ser.get("name"), len(ser["values"]), len(op["categories"])))
                    data.add_series(ser.get("name", ""), ser["values"])
                shape.chart.replace_data(data)
                log.append("replaced chart data in '%s' on slide id %d" % (shape.name, slide.slide_id))
            elif kind == "delete_shape":
                slide = resolve_slide(prs, op["slide"], original)
                shape = resolve_shape(slide, op["shape"])
                shape._element.getparent().remove(shape._element)
                log.append("deleted shape '%s' on slide id %d" % (shape.name, slide.slide_id))
            elif kind == "move_slide":
                slide = resolve_slide(prs, op["slide"], original)
                after = op.get("after", 0)
                if isinstance(after, dict):
                    after = original.get(int(after["n"]))
                move_after(prs, slide, after)
                log.append("moved slide id %d" % slide.slide_id)
            elif kind == "delete_slides":
                for ref in op["slides"]:
                    slide = resolve_slide(prs, ref, original)
                    sid = slide.slide_id
                    delete_slide(prs, slide)
                    log.append("deleted slide id %d" % sid)
            elif kind == "insert_slides":
                made = insert_slides(prs, info, op, original, findings)
                log.append("inserted %d new slide(s): ids %s" % (len(made), ", ".join(map(str, made))))
            elif kind == "replace_slide":
                slide = resolve_slide(prs, op["slide"], original)
                sid = slide.slide_id
                made = insert_slides(prs, info, {"after": sid, "slides": [op["with"]], "theme": op.get("theme")}, original, findings)
                delete_slide(prs, slide)
                log.append("replaced slide id %d with new slide id %s" % (sid, made[0]))
            else:
                raise EditError("Unknown op %r. Ops: set_text, replace_text, set_notes, set_table_cell, set_chart_data, delete_shape, move_slide, delete_slides, insert_slides, replace_slide." % kind)
        except EditError as exc:
            findings.append({"slide": 0, "level": "error", "msg": "Edit %d (%s): %s" % (k, kind, exc)})
        except (KeyError, ValueError, IndexError) as exc:
            findings.append({"slide": 0, "level": "error", "msg": "Edit %d (%s): missing or invalid field %s" % (k, kind, exc)})
    return log


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("command", choices=["inspect", "apply"])
    ap.add_argument("deck")
    ap.add_argument("edits", nargs="?")
    ap.add_argument("-o", "--out")
    args = ap.parse_args()
    try:
        prs = Presentation(args.deck)
    except Exception as exc:
        print("Could not open %s: %s" % (args.deck, exc))
        return 2
    if args.command == "inspect":
        info = inspect(prs, args.deck)
        out = args.out or os.path.splitext(args.deck)[0] + ".outline.json"
        with open(out, "w", encoding="utf-8") as handle:
            json.dump(info, handle, indent=1, ensure_ascii=False, default=str)
        print_outline(info)
        print("Full outline: %s" % out)
        return 0
    if not args.edits or not args.out:
        print("apply needs an edits.json file and -o output.pptx")
        return 2
    with open(args.edits, encoding="utf-8") as handle:
        ops = json.load(handle)
    ops = ops.get("edits", ops) if isinstance(ops, dict) else ops
    findings = []
    log = apply(prs, ops, findings)
    renumber_parts(prs)
    prs.save(args.out)
    errors = sum(1 for f in findings if f["level"] == "error")
    print("Saved %s · %d edit(s) applied" % (args.out, len(log)))
    for line in log:
        print("  - " + line)
    if findings:
        print(dk_lint.format_findings(findings, limit=60))
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
