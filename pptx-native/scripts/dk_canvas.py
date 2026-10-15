"""Deck and Canvas: the drawing primitives every layout uses.

All coordinates are inches on a 13.333 x 7.5 slide. Every primitive creates a native,
editable PowerPoint object and records its box so dk_lint can check bounds, overflow and contrast.
"""

import json
import os
import re
import uuid

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from dk_units import Inches, Pt, emu

import dk_images
import dk_measure as M
from dk_theme import Theme, contrast, mix, pick_on, readable

W, H = 13.333, 7.5
MX = 0.75            # left and right margin
TOP = 0.62           # top margin
FOOT_Y = 6.93        # footer text baseline box top
BOTTOM = 6.72        # lowest y for body content when a footer is present
CW = W - 2 * MX      # content width
LH = 1.2             # PowerPoint line pitch at spacing 1.0, in ems

SHAPES = {
    "rect": MSO_SHAPE.RECTANGLE, "round": MSO_SHAPE.ROUNDED_RECTANGLE, "oval": MSO_SHAPE.OVAL,
    "chevron": MSO_SHAPE.CHEVRON, "pentagon": MSO_SHAPE.PENTAGON, "arrow": MSO_SHAPE.RIGHT_ARROW,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE, "diamond": MSO_SHAPE.DIAMOND,
    "down-arrow": MSO_SHAPE.DOWN_ARROW, "hexagon": MSO_SHAPE.HEXAGON,
}
ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}
ICON_ALIAS = {
    "home": "building-2", "bar-chart": "chart-column", "line-chart": "chart-line", "pie-chart": "chart-pie",
    "check-circle": "circle-check", "help": "circle-help", "help-circle": "circle-help", "question": "circle-help",
    "warning": "triangle-alert", "alert": "triangle-alert", "money": "banknote", "dollar": "dollar-sign",
    "rupee": "indian-rupee", "people": "users", "team": "users", "person": "user", "idea": "lightbulb",
    "growth": "trending-up", "decline": "trending-down", "time": "clock", "security": "shield-check",
    "ai": "sparkles", "data": "database", "goal": "target", "launch": "rocket", "settings": "settings",
    "doc": "file-text", "document": "file-text", "email": "mail", "chat": "message-circle", "phone": "smartphone",
    "world": "globe", "location": "map-pin", "cart": "shopping-cart", "risk": "triangle-alert", "tools": "wrench",
    "process": "workflow", "speed": "gauge", "quality": "badge-check", "cost": "coins", "win": "trophy",
}
_icons = None


def _c(hexstr):
    return RGBColor.from_string(hexstr)


def _plain_style(shape):
    """Drop the theme style reference so no renderer adds a shadow, fill or outline we did not set."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


DECORATIVE_URI = "{C183D7F6-B498-43B3-948B-1728B52AA6E4}"
DECORATIVE_NS = "http://schemas.microsoft.com/office/drawing/2017/decorative"
SECTIONS_URI = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"
P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"


def _cnvpr(element):
    """The cNvPr (id, name, alt text) of a shape, connector, group, picture or graphic frame."""
    found = element.xpath("./*[1]/p:cNvPr")
    return found[0] if found else None


def set_alt(shape, text):
    """Alt text that screen readers and PowerPoint's Accessibility Checker read."""
    props = _cnvpr(shape._element)
    if props is not None and text:
        props.set("descr", str(text)[:1000])


def mark_decorative(element):
    """Flag a purely visual object as decorative, so screen readers skip it and the checker does not ask for alt text."""
    props = _cnvpr(element)
    if props is None:
        return
    ext_list = props.find(qn("a:extLst"))
    if ext_list is None:
        ext_list = props.makeelement(qn("a:extLst"), {})
        props.append(ext_list)
    for ext in ext_list.findall(qn("a:ext")):
        if ext.get("uri") == DECORATIVE_URI:
            return
    ext = ext_list.makeelement(qn("a:ext"), {"uri": DECORATIVE_URI})
    flag = ext.makeelement("{%s}decorative" % DECORATIVE_NS, {"val": "1"}, nsmap={"adec": DECORATIVE_NS})
    ext.append(flag)
    ext_list.append(ext)


def make_title(shape):
    """Turn a text box into the slide's title placeholder. Its look is unchanged (every format is set
    on the text itself), but PowerPoint, Keynote and screen readers now see it as the slide title."""
    sp = shape._element
    c_sp = sp.find(qn("p:nvSpPr")).find(qn("p:cNvSpPr"))
    if c_sp is not None and c_sp.get("txBox"):
        del c_sp.attrib["txBox"]
    nv = sp.find(qn("p:nvSpPr")).find(qn("p:nvPr"))
    for old in nv.findall(qn("p:ph")):
        nv.remove(old)
    nv.insert(0, nv.makeelement(qn("p:ph"), {"type": "title"}))


def load_icons():
    global _icons
    if _icons is None:
        path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "icons.json")
        try:
            with open(path) as handle:
                _icons = json.load(handle)["icons"]
        except Exception:
            _icons = {}
    return _icons


def icon_names():
    return sorted(load_icons())


def parse_runs(text):
    """Split '**bold** and plain' into runs. Returns [(text, bold)]."""
    parts = re.split(r"(\*\*[^*]+\*\*)", str(text))
    return [(p[2:-2], True) if p.startswith("**") and p.endswith("**") else (p, False) for p in parts if p]


def plain(text):
    return re.sub(r"\*\*([^*]+)\*\*", r"\1", str(text))


def blank_layout(prs):
    """The layout with the fewest placeholders, preferring one named Blank."""
    layouts = list(prs.slide_layouts)
    named = [l for l in layouts if "blank" in (l.name or "").lower()]
    pool = named or layouts
    return min(pool, key=lambda l: len(l.placeholders))


class Deck(object):
    def __init__(self, theme=None, meta=None, prs=None, template_marks=None):
        self.theme = theme if isinstance(theme, Theme) else Theme(theme)
        self.meta = dict(meta or {})
        if prs is None:
            self.prs = Presentation()
            self.prs.slide_width = Inches(W)
            self.prs.slide_height = Inches(H)
            self.blank = self.prs.slide_layouts[6]
        else:
            self.prs = prs  # an existing deck: new slides use its blank layout and size
            self.blank = blank_layout(prs)
        self.canvases = []
        # a template's fixed graphics (logo, bars): light slides let them show, and text keeps clear of them
        self.marks = list(template_marks or [])
        self.template = template_marks is not None
        # an existing deck keeps its own sections (edit mode); new decks, templates included, get one per divider
        self.keep_sections = prs is not None and not self.template
        props = self.prs.core_properties
        props.title = self.meta.get("title", "")
        props.author = self.meta.get("author", "")
        props.subject = self.meta.get("subtitle", "")

    def new_slide(self, tone="light", spec=None):
        slide = self.prs.slides.add_slide(self.blank)
        for ph in list(slide.placeholders):  # empty placeholders would show "Click to add" boxes
            ph._element.getparent().remove(ph._element)
        canvas = Canvas(self, slide, tone, len(self.canvases) + 1, spec or {})
        self.canvases.append(canvas)
        return canvas

    def save(self, path):
        for canvas in self.canvases:
            canvas.finish()
        self.write_sections()
        self.prs.save(path)
        return path

    def write_sections(self):
        """Section dividers become PowerPoint sections, so long decks can be collapsed and navigated by part."""
        if not any(c.layout == "section" for c in self.canvases) or self.keep_sections:
            return
        groups, name = [], "Opening"
        for canvas in self.canvases:
            if canvas.layout == "section":
                name = plain(canvas.spec.get("title") or "Section %d" % (len(groups) + 1))[:60]
                groups.append((name, []))
            elif not groups:
                groups.append((name, []))
            groups[-1][1].append(canvas.slide.slide_id)
        root = self.prs.part._element
        ext_list = root.find(qn("p:extLst"))
        if ext_list is None:
            ext_list = root.makeelement(qn("p:extLst"), {})
            root.append(ext_list)
        for ext in ext_list.findall(qn("p:ext")):
            if ext.get("uri") == SECTIONS_URI:
                ext_list.remove(ext)
        ext = ext_list.makeelement(qn("p:ext"), {"uri": SECTIONS_URI})
        tag = "{%s}" % P14_NS
        lst = ext.makeelement(tag + "sectionLst", {}, nsmap={"p14": P14_NS})
        for title, ids in groups:
            sec = lst.makeelement(tag + "section", {"name": title, "id": "{%s}" % str(uuid.uuid4()).upper()})
            id_list = sec.makeelement(tag + "sldIdLst", {})
            for sid in ids:
                id_list.append(id_list.makeelement(tag + "sldId", {"id": str(sid)}))
            sec.append(id_list)
            lst.append(sec)
        ext.append(lst)
        ext_list.append(ext)


class Canvas(object):
    def __init__(self, deck, slide, tone, number, spec):
        self.deck, self.slide, self.number, self.spec = deck, slide, number, spec
        self.theme = deck.theme
        self.P = self.theme.tone(tone)
        self.boxes, self.warnings, self.errors = [], [], []
        self.layout = spec.get("layout", "")
        self.bottom = BOTTOM
        self.titled = False
        self.finished = False
        self.body_from = (0, 0)  # (first shape, first box) of the slide body, set after the header
        self.fill_ratio = None
        if not (deck.template and tone == "light"):  # light slides in a template show the master's own background
            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = _c(self.P["bg"])

    # ---- accessibility -----------------------------------------------------------------------
    def finish(self):
        """Last pass before saving: a real title on every slide, alt text where it matters,
        and decorative marks on purely visual shapes."""
        if self.finished:
            return
        self.finished = True
        if not self.titled:
            hidden = plain(self.spec.get("title") or self.spec.get("kicker") or self.spec.get("text")
                           or self.spec.get("quote") or self.deck.meta.get("title") or "Slide %d" % self.number)
            box = self.slide.shapes.add_textbox(Inches(MX), Inches(-1.2), Inches(CW), Inches(0.6))
            box.name = "Title (off-slide)"
            label = hidden[:200]
            frame = box.text_frame
            frame.text = label
            make_title(box)
            box._element.getparent().remove(box._element)
            self.slide.shapes._spTree.insert(2, box._element)  # first in reading order
            self.titled = True
        for shape in self.slide.shapes:
            el = shape._element
            tag = el.tag.split("}")[-1]
            if tag == "grpSp" or tag == "cxnSp":
                mark_decorative(el)
            elif tag == "sp" and not (shape.has_text_frame and shape.text_frame.text.strip()):
                mark_decorative(el)
        # Reading order follows drawing order, which is already title, body, footer. Shapes are never
        # reordered here, because order is also stacking order and moving them would hide things.

    # ---- balance -----------------------------------------------------------------------------
    def mark_body(self):
        """Everything drawn after this point is the slide body (used by settle)."""
        self.body_from = (len(self.slide.shapes._spTree) - 2, len(self.boxes))

    def settle(self, top, bottom, bias=0.42, max_shift=0.5, min_free=0.4):
        """Ease the body down toward the optical centre between top and bottom when it would otherwise
        hug the headline over an empty lower slide. The shift is capped (half an inch by default) so the
        line where content starts stays consistent from slide to slide; layouts grow first, settle second.
        The takeaway stays put. Returns the shift in inches."""
        recs = [r for r in self.boxes[self.body_from[1]:] if r.get("name") != "Takeaway"]
        if not recs:
            return 0.0
        used_top = min(r["y"] for r in recs)
        used_bottom = max(r["y"] + r["h"] for r in recs)
        span = bottom - top
        if span > 0:
            self.fill_ratio = (used_bottom - used_top) / span
        free = (bottom - used_bottom) - max(0.0, used_top - top)
        if free < min_free:
            return 0.0
        shift = min(max_shift, free * bias)
        for shape in list(self.slide.shapes)[self.body_from[0]:]:
            if shape.name == "Takeaway" or shape.top is None:
                continue
            shape.top = shape.top + emu(shift)
        for r in recs:
            r["y"] += shift
        return shift

    # ---- bookkeeping -------------------------------------------------------------------------
    def warn(self, message):
        self.warnings.append(message)

    def _rec(self, kind, x, y, w, h, **extra):
        rec = {"kind": kind, "x": x, "y": y, "w": w, "h": h}
        rec.update(extra)
        self.boxes.append(rec)
        return rec

    def font(self, role):
        return self.theme.get(role, role) if role in ("display", "body", "mono") else role

    def display_bold(self):
        return self.theme["display"] != "Arial Black"

    def radius(self, scale=1.0):
        return self.theme["radius"] * scale

    # ---- shapes ------------------------------------------------------------------------------
    def rect(self, x, y, w, h, fill=None, line=None, line_w=1.0, radius=None, kind=None, alpha=None,
             name=None, rotation=0, dash=False):
        """A rectangle (rounded when radius > 0) or another SHAPES kind. fill and line are hex or None."""
        if w <= 0 or h <= 0:
            return None
        radius = self.theme["radius"] if radius is None else radius
        if kind is None:
            kind = "round" if radius > 0.005 else "rect"
        shape = self.slide.shapes.add_shape(SHAPES[kind], Inches(x), Inches(y), Inches(w), Inches(h))
        if kind == "round":
            try:
                shape.adjustments[0] = max(0.0, min(0.5, radius / min(w, h)))
            except Exception:
                pass
        if fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = _c(fill)
            if alpha is not None:
                clr = shape.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
                a = clr.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))})
                clr.append(a)
        else:
            shape.fill.background()
        if line:
            shape.line.color.rgb = _c(line)
            shape.line.width = Pt(line_w)
            if dash:
                from pptx.enum.dml import MSO_LINE
                shape.line.dash_style = MSO_LINE.DASH
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False
        _plain_style(shape)
        if rotation:
            shape.rotation = rotation
        if name:
            shape.name = name
        shape.text_frame.text = ""
        self._rec("shape", x, y, w, h, name=name or kind, fill=fill)
        return shape

    def oval(self, x, y, w, h, fill=None, line=None, line_w=1.0, name=None, alpha=None):
        return self.rect(x, y, w, h, fill=fill, line=line, line_w=line_w, kind="oval", name=name, alpha=alpha)

    def line(self, x1, y1, x2, y2, color=None, w=1.0, dash=False, arrow=False, name=None):
        conn = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        conn.line.color.rgb = _c(color or self.P["line"])
        conn.line.width = Pt(w)
        _plain_style(conn)
        if dash:
            from pptx.enum.dml import MSO_LINE
            conn.line.dash_style = MSO_LINE.DASH
        if arrow:
            ln = conn.line._get_or_add_ln()
            ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
        if name:
            conn.name = name
        return conn

    # ---- text --------------------------------------------------------------------------------
    def _paragraphs(self, content):
        if isinstance(content, (list, tuple)):
            return [str(p) for p in content]
        return str(content).split("\n")

    def measure(self, content, size, w, font="body", bold=False, spacing=1.0, para_gap=0, tracking=0, caps=False):
        paras = [plain(p).upper() if caps else plain(p) for p in self._paragraphs(content)]
        return M.text_height(paras, self.font(font), size, w, bold, spacing * LH, para_gap, self._spc(size, tracking))

    def _spc(self, size, tracking):
        return int(round(size * tracking / 10.0)) if tracking else 0

    def text(self, x, y, w, h, content, size=14, color=None, font="body", bold=False, italic=False,
             align="l", valign="t", spacing=1.0, tracking=0, caps=False, on=None, fit=None,
             max_lines=None, para_gap=0, name=None, bullets=False, role="text", rotation=0):
        """A text box at a fixed position. fit=min_size shrinks text until it fits the box.

        content is a string ('\\n' separates paragraphs, **bold** marks bold runs) or a list of paragraphs.
        Returns the font size used.
        """
        paras = self._paragraphs(content)
        family = self.font(font)
        color = color or self.P["ink"]
        measured = [plain(p).upper() if caps else plain(p) for p in paras]
        indent = 0.22 if bullets else 0.0
        fits = True
        if fit:
            size, fits = M.fit_size(measured, family, w - indent, h, size, fit, bold, spacing * LH, para_gap,
                                    self._spc(size, tracking), max_lines)
        need, lines = M.text_height(measured, family, size, w - indent, bold, spacing * LH, para_gap,
                                    self._spc(size, tracking))
        if M.longest_word(measured, family, size, bold, self._spc(size, tracking)) > w - indent + 0.01:
            fits = False
        box = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(max(h, 0.05)))
        if name:
            box.name = name
        if rotation:
            box.rotation = rotation
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = ANCHOR[valign]
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = ALIGN[align]
            p.line_spacing = spacing
            if para_gap and i < len(paras) - 1:
                p.space_after = Pt(para_gap)
            if bullets:
                self._bullet(p, color, indent)
            for piece, is_bold in parse_runs(para):
                run = p.add_run()
                label = piece.upper() if caps else piece
                run.text = label
                f = run.font
                f.name = family
                f.size = Pt(size)
                f.bold = bool(bold or is_bold)
                f.italic = bool(italic)
                f.color.rgb = _c(color)
                spc = self._spc(size, tracking)
                if spc:
                    run._r.get_or_add_rPr().set("spc", str(spc))
        if role == "title" and not self.titled:
            make_title(box)
            self.titled = True
            # read the title first, unless moving it to the back would put it under a panel or picture
            earlier = [r for r in self.boxes if r["kind"] in ("shape", "image")]
            if not any(r["x"] < x + w and x < r["x"] + r["w"] and r["y"] < y + h and y < r["y"] + r["h"] for r in earlier):
                box._element.getparent().remove(box._element)
                self.slide.shapes._spTree.insert(2, box._element)
        if rotation in (90, -90, 270):  # record the box as it sits on the slide after rotation
            x, y, w, h, need = x + w / 2 - h / 2, y + h / 2 - w / 2, h, w, 0
        self._rec("text", x, y, w, h, need=need, size=size, lines=lines, text=" / ".join(measured)[:90],
                  color=color, on=on or self.P["bg"], fits=fits, name=name or role, role=role)
        return size

    def _bullet(self, p, color, indent):
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(indent))))
        pPr.set("indent", str(-int(Inches(indent))))
        for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone"):
            for el in pPr.findall(qn(tag)):
                pPr.remove(el)
        clr = pPr.makeelement(qn("a:buClr"), {})
        clr.append(clr.makeelement(qn("a:srgbClr"), {"val": color}))
        pPr.append(clr)
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "\u2022"}))

    def balance(self, content, size, w, font="display", bold=True, tracking=0):
        """Narrowest width that keeps the same line count, so a headline does not leave one word alone."""
        fam, spc = self.font(font), self._spc(size, tracking)
        text = plain(content)
        lines = M.wrap(text, fam, size, w, bold, spc)
        if len(lines) < 2 or "\n" in text:
            return w
        best = w
        trial = w
        while trial > w * 0.55:
            trial -= 0.2
            if len(M.wrap(text, fam, size, trial, bold, spc)) > len(lines) or M.longest_word([text], fam, size, bold, spc) > trial:
                break
            best = trial
        return best

    def put(self, x, y, w, content, size=14, max_h=None, min_size=None, gap=0.0, **kw):
        """Auto-height text: the box is as tall as the text needs. Returns the y just below it (plus gap).

        With max_h, text shrinks toward min_size to stay inside; if it still overflows, lint reports it.
        """
        if content is None or content == "" or content == []:
            return y
        spacing, para_gap = kw.get("spacing", 1.0), kw.get("para_gap", 0)
        bullets = kw.get("bullets", False)
        inner = w - (0.22 if bullets else 0.0)
        kw.pop("fit", None)
        if max_h is not None and min_size:
            paras = [plain(p).upper() if kw.get("caps") else plain(p) for p in self._paragraphs(content)]
            size, _ = M.fit_size(paras, self.font(kw.get("font", "body")), inner, max_h, size, min_size,
                                 kw.get("bold", False), spacing * LH, para_gap,
                                 self._spc(size, kw.get("tracking", 0)), kw.get("max_lines"))
        need, _ = self.measure(content, size, inner, kw.get("font", "body"), kw.get("bold", False), spacing,
                               para_gap, kw.get("tracking", 0), kw.get("caps", False))
        h = need if max_h is None else min(need, max_h)
        kw.pop("max_lines", None)
        if kw.get("role") == "title" and isinstance(content, str) and kw.get("align", "l") == "l":
            w = self.balance(content, size, w, kw.get("font", "body"), kw.get("bold", False), kw.get("tracking", 0)) + 0.05
        self.text(x, y, w, h, content, size=size, **kw)
        if max_h is not None and need > max_h + 0.02:
            self.boxes[-1]["need"] = need
            self.boxes[-1]["fits"] = False
        return y + h + gap

    def pill(self, x, y, label, fill=None, color=None, size=10, pad=0.14, h=None, bold=True, font="body", line=None):
        """A small rounded label sized to its text. Returns its width."""
        fill = fill or self.P["accent"]
        color = color or pick_on(fill, "141414", "FFFFFF")
        h = h or size / 72.0 * 2.1
        w = M.text_width(label.upper(), self.font(font), size, bold, self._spc(size, 40)) + 2 * pad
        self.rect(x, y, w, h, fill=fill, line=line, radius=min(h / 2, max(self.theme["radius"], 0.02) * 1.6), name="Pill")
        self.text(x, y, w, h, label, size=size, color=color, bold=bold, align="c", valign="m", caps=True,
                  tracking=40, on=fill, font=font, role="pill")
        return w

    def badge(self, x, y, d, label, fill=None, color=None, size=None, square=False, font="display"):
        """A circle (or rounded square) with a short label such as a number or letter."""
        fill = fill or self.P["ink"]
        color = color or pick_on(fill, "141414", "FFFFFF")
        if square:
            self.rect(x, y, d, d, fill=fill, radius=min(d * 0.28, max(self.theme["radius"], 0.03) * 1.4), name="Badge")
        else:
            self.oval(x, y, d, d, fill=fill, name="Badge")
        self.text(x, y, d, d, str(label), size=size or d * 72 * 0.42, color=color, bold=True, align="c",
                  valign="m", on=fill, font=font, role="badge")

    # ---- pictures ----------------------------------------------------------------------------
    def picture(self, image, x, y, w, h, fit=None, focus=None, radius=None, name="Picture", frame=False):
        """A user-supplied picture. image is a path or {"path", "alt", "fit", "focus", "decorative"}.

        fit "cover" fills the box and crops the overflow (photos); "contain" shows the whole image
        (screenshots, charts, logos). Returns the box the picture occupies, or None if it failed.
        """
        spec = image if isinstance(image, dict) else {"path": image}
        fit = spec.get("fit") or fit or "cover"
        focus = spec.get("focus") or focus or "center"
        info = dk_images.load(spec.get("path") or spec.get("src") or "")
        if info.get("error"):
            self.errors.append(info["error"])
            self.rect(x, y, w, h, fill=mix(self.P["bg"], self.P["ink"], 0.06), line=self.P["line"], dash=True, name="Missing picture")
            self.text(x + 0.2, y, w - 0.4, h, "Missing picture: " + str(spec.get("path") or "")[:60], size=11,
                      color=self.P["muted"], align="c", valign="m", fit=8, name="Missing picture label")
            return None
        if frame:  # screenshots sit on a quiet panel so their edges are visible on any background
            self.rect(x, y, w, h, fill=self.P["surface"], line=self.P["line"], name="Picture frame")
            pad = 0.12
            x, y, w, h = x + pad, y + pad, w - 2 * pad, h - 2 * pad
        if fit == "contain":
            x, y, w, h = dk_images.contain_box(info["px"], x, y, w, h, spec.get("align") or "center")
            crop = (0.0, 0.0, 0.0, 0.0)
        else:
            crop = dk_images.cover_crop(info["px"], w, h, focus)
        pic = self.slide.shapes.add_picture(info["path"], Inches(x), Inches(y), Inches(w), Inches(h))
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
        pic.name = spec.get("name") or name
        radius = self.theme["radius"] if radius is None else radius
        if radius > 0.005 and not frame:
            geom = pic._element.spPr.find(qn("a:prstGeom"))
            if geom is not None:
                geom.set("prst", "roundRect")
                av = geom.find(qn("a:avLst"))
                if av is None:
                    av = geom.makeelement(qn("a:avLst"), {})
                    geom.append(av)
                for old in list(av):
                    av.remove(old)
                adj = int(min(50000, radius / min(w, h) * 100000))
                av.append(av.makeelement(qn("a:gd"), {"name": "adj", "fmla": "val %d" % adj}))
        if spec.get("decorative"):
            mark_decorative(pic._element)
        elif spec.get("alt"):
            set_alt(pic, spec["alt"])
        else:
            self.warn("Picture %s has no alt text. Add \"alt\" (what it shows, in a sentence) or \"decorative\": true." % info["name"])
        shown_px = info["px"][0] * (1 - crop[0] - crop[2])
        ppi = shown_px / max(w, 0.01)
        if ppi < 100:
            self.warn("Picture %s is low resolution at this size (%d pixels per inch); use a larger file or show it smaller." % (info["name"], ppi))
        self._rec("image", x, y, w, h, name=pic.name)
        return x, y, w, h

    # ---- icons -------------------------------------------------------------------------------
    def icon(self, name, x, y, size, color=None, weight=None):
        """A native vector icon (grouped freeform lines) from the bundled Lucide subset."""
        icons = load_icons()
        key = str(name or "").strip().lower().replace("_", "-").replace(" ", "-")
        key = key if key in icons else ICON_ALIAS.get(key, key)
        color = color or self.P["ink"]
        if key not in icons:
            self.warn("Unknown icon %r; drew a dot. Run build_deck.py --icons for the list." % name)
            self.oval(x + size * 0.35, y + size * 0.35, size * 0.3, size * 0.3, fill=color)
            return None
        weight = weight or max(1.0, size * 72 / 24.0 * 0.085 * 24 / 24.0 * 1.0)
        unit = size / 24.0
        group = self.slide.shapes.add_group_shape()
        group.name = "Icon " + key
        for prim in icons[key]:
            if prim[0] == "c":
                _, cx, cy, r = prim
                r = max(r, 0.45)
                shp = group.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + (cx - r) * unit), Inches(y + (cy - r) * unit),
                                             Inches(2 * r * unit), Inches(2 * r * unit))
                shp.text_frame.text = ""
            else:
                _, closed, pts = prim
                emu = [(int(Inches(x + px * unit)), int(Inches(y + py * unit))) for px, py in pts]
                fb = group.shapes.build_freeform(emu[0][0], emu[0][1])
                fb.add_line_segments(emu[1:], close=bool(closed))
                shp = fb.convert_to_shape()
            shp.fill.background()
            shp.line.color.rgb = _c(color)
            shp.line.width = Pt(weight)
            shp.shadow.inherit = False
            _plain_style(shp)
            ln = shp.line._get_or_add_ln()
            ln.set("cap", "rnd")
            ln.append(ln.makeelement(qn("a:round"), {}))
        self._rec("icon", x, y, size, size, name=key)
        return group

    def icon_badge(self, name, x, y, d, fill=None, color=None, round_=False):
        """An icon centred on a filled tile."""
        fill = fill or self.P["ink"]
        color = color or pick_on(fill, "141414", "FFFFFF")
        if round_:
            self.oval(x, y, d, d, fill=fill, name="Icon tile")
        else:
            self.rect(x, y, d, d, fill=fill, radius=min(d * 0.26, max(self.theme["radius"], 0.03) * 1.5), name="Icon tile")
        pad = d * 0.24
        self.icon(name, x + pad, y + pad, d - 2 * pad, color=color)

    # ---- composite pieces --------------------------------------------------------------------
    def card(self, x, y, w, h, style=None, tab=None, emphasis=False, pad=None, fill=None, name="Card"):
        """A content card in the theme's card style. Returns (ix, iy, iw, ih, fill, ink, muted)."""
        P = self.P
        style = style or self.theme["card"]
        pad = 0.26 if pad is None else pad
        ink, muted = P["ink"], P["muted"]
        if emphasis:
            fill = fill or P["ink"]
            ink = pick_on(fill, "141414", "FFFFFF")
            muted = mix(ink, fill, 0.28)
            self.rect(x, y, w, h, fill=fill, name=name)
        elif style == "fill":
            fill = fill or P["surface"]
            self.rect(x, y, w, h, fill=fill, name=name)
        elif style == "bold":
            fill = fill or P["surface"]
            self.rect(x, y, w, h, fill=fill, line=P["ink"], line_w=1.75, name=name)
        elif style == "rule":
            fill = fill or P["bg"]
            self.rect(x, y, w, 0.025, fill=P["ink"], radius=0, name=name + " rule")
            self._rec("region", x, y, w, h, name=name)
            x, w, pad = x, w, 0.0
            y += 0.2
            h -= 0.2
        else:
            fill = fill or P["surface"]
            self.rect(x, y, w, h, fill=fill, line=P["line"], line_w=1.0, name=name)
        if fill != P["bg"] and not emphasis:
            muted = readable(muted, fill, 4.5)
        top = y + pad
        if tab:
            self.rect(x + pad, top, 0.6, 0.08, fill=tab, radius=0.0, name=name + " tab")
            top += 0.24
        return x + pad, top, w - 2 * pad, (y + h - pad) - top, fill, ink, muted

    def header(self, kicker=None, title=None, subtitle=None, width=None, size=None, max_lines=2):
        """Slide header in the theme's header style. Returns the y where body content starts."""
        P, T = self.P, self.theme
        style = T["header"]
        x, y = MX, TOP
        width = width or CW
        if style == "block" and title:
            x += 0.26
            width -= 0.26
        if kicker:
            y = self.kicker(x, y, kicker) + 0.1
        y_title = y
        if title:
            size = size or T.get("title_size", 30)
            y = self.put(x, y, width, title, size=size, max_h=size * LH * 0.98 * max_lines / 72.0 + 0.02,
                         min_size=size - 7, font="display", bold=self.display_bold(), spacing=0.98,
                         tracking=T["tracking"], max_lines=max_lines, name="Title", role="title")
            if style == "block":
                self.rect(MX, y_title + 0.04, 0.09, y - y_title - 0.08, fill=P["accent"], radius=0, name="Title block")
        if subtitle:
            y = self.put(x, y + 0.1, min(width, 10.6), subtitle, size=16, color=P["muted"], max_h=0.62, min_size=13,
                         spacing=1.05, name="Subtitle", role="subtitle")
        if style == "rule" and title:
            y += 0.16
            self.rect(MX, y, CW, 0.02, fill=P["ink"] if T["card"] == "rule" else P["line"], radius=0, name="Header rule")
            y += 0.02
        self.mark_body()
        return y + 0.3

    def kicker(self, x, y, label, color=None):
        """The small label above a headline. Returns the y below it."""
        P, style = self.P, self.theme["kicker"]
        if style == "pill":
            h = 0.3
            self.pill(x, y, label, fill=P["accent"], size=10, h=h)
            return y + h
        if style == "mono":
            return self.put(x, y, CW, label, size=11, color=color or P["muted"], font="mono", caps=True,
                            tracking=30, name="Kicker", role="kicker")
        return self.put(x, y, CW, label, size=11, color=color or P["muted"], bold=True, caps=True, tracking=90,
                        name="Kicker", role="kicker")

    def footer(self, source=None, page=True):
        """Source line (left) and deck label with page number (right)."""
        P = self.P
        label = self.deck.meta.get("footer", "")
        right = ""
        if page and self.deck.meta.get("page_numbers", True):
            right = ("%s   %d" % (label, self.number)).strip() if label else str(self.number)
        rw = 0.0
        edge = W - MX
        for m in self.deck.marks:  # master graphics show on every slide, whatever its tone
            if m["y"] < FOOT_Y + 0.25 and m["y"] + m["h"] > FOOT_Y and m["x"] + m["w"] > W / 2:
                edge = min(edge, m["x"] - 0.2)
        if right:
            rw = min(4.0, M.text_width(right, self.font("body"), 9) + 0.1)
            self.text(edge - rw, FOOT_Y, rw, 0.2, right, size=9, color=P["muted"], align="r", name="Page", role="footer")
            rw += W - MX - edge
        if source:
            self.put(MX, FOOT_Y, CW - rw - 0.3, source, size=9, color=P["muted"], max_h=0.36, min_size=8,
                     name="Source", role="footer")

    def notes(self, text):
        if text:
            self.slide.notes_slide.notes_text_frame.text = str(text)

    def stat(self, x, y, w, value, label=None, size=40, color=None, on=None, label_color=None, align="l", max_h=None):
        """A big number with a small label under it. Returns the y below."""
        P = self.P
        y = self.put(x, y, w, str(value), size=size, max_h=size * LH / 72.0 * 1.02, min_size=size * 0.55,
                     font="display", bold=self.display_bold(), color=color or P["ink"], spacing=0.95,
                     tracking=self.theme["tracking"], max_lines=1, on=on, align=align, name="Stat", role="stat")
        if label:
            y = self.put(x, y + 0.04, w, label, size=12, color=label_color or P["muted"], on=on, align=align,
                         max_h=max_h, min_size=10, spacing=1.05, name="Stat label")
        return y
