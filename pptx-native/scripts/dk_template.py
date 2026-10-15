"""Build a new deck inside the user's own template (.pptx or .potx).

The template's masters, layouts, theme colours, fonts and fixed graphics (logo, brand bars) are kept.
Its sample slides are removed. New slides use its blank layout, so the master's background and
graphics show through, and the kit takes its palette and fonts from the template's theme.
"""

import os
import shutil
import tempfile
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu

DESIGN_W = 13.333
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
POTX = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"


class TemplateError(Exception):
    pass


def _as_pptx(path):
    """python-pptx opens .pptx only; a .potx differs in one content type, so copy and relabel it."""
    if not path.lower().endswith((".potx", ".potm")):
        return path
    out = os.path.join(tempfile.mkdtemp(prefix="pptx-native-tpl-"), "template.pptx")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(POTX.encode(), PPTX.encode())
            dst.writestr(item, data)
    return out


def theme_colors_fonts(prs):
    colors, fonts = {}, {}
    part = prs.slide_masters[0].part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    scheme = root.find(".//%sclrScheme" % A)
    for el in (scheme if scheme is not None else []):
        srgb = el.find("%ssrgbClr" % A)
        sysc = el.find("%ssysClr" % A)
        val = srgb.get("val") if srgb is not None else (sysc.get("lastClr") if sysc is not None else None)
        if val:
            colors[etree.QName(el).localname] = val.upper()
    for key, tag in (("major", "majorFont"), ("minor", "minorFont")):
        latin = root.find(".//%s%s/%slatin" % (A, tag, A))
        if latin is not None and latin.get("typeface") and not latin.get("typeface").startswith("+"):
            fonts[key] = latin.get("typeface")
    return colors, fonts


def theme_tokens(colors, fonts):
    accents = [colors[k] for k in ("accent1", "accent2", "accent3", "accent4") if k in colors]
    t = {"bg": colors.get("lt1", "FFFFFF"), "ink": colors.get("dk1", "141414"), "dark": colors.get("dk2", colors.get("dk1", "141414")),
         "accents": accents or None,
         "series": [colors[k] for k in ("accent1", "accent2", "accent3", "accent4", "accent5", "accent6") if k in colors] or None,
         "display": fonts.get("major"), "body": fonts.get("minor")}
    return {k: v for k, v in t.items() if v}


def blank_layout(prs):
    layouts = list(prs.slide_layouts)
    named = [l for l in layouts if "blank" in (l.name or "").lower()]
    return min(named or layouts, key=lambda l: len(l.placeholders))


def marks(prs, layout, scale):
    """Boxes (design inches) of the fixed graphics a slide on this layout inherits: logos, bars, art."""
    out = []
    for owner in (prs.slide_masters[0], layout):
        for shape in owner.shapes:
            if shape.is_placeholder or shape.width is None or shape.height is None:
                continue
            w, h = Emu(shape.width).inches / scale, Emu(shape.height).inches / scale
            if w * h < 0.01:
                continue
            out.append({"x": Emu(shape.left).inches / scale, "y": Emu(shape.top).inches / scale, "w": w, "h": h,
                        "name": shape.name or "template graphic"})
    return out


def open_template(path):
    """Returns (presentation with no slides, theme tokens, design scale, marks, notes)."""
    if not os.path.exists(path):
        raise TemplateError("Template not found: %s. Use the attached file's exact name." % path)
    try:
        prs = Presentation(_as_pptx(path))
    except Exception as exc:
        raise TemplateError("Cannot open the template %s (%s). Save it as .pptx or .potx and attach it again." % (os.path.basename(path), type(exc).__name__))
    w, h = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    if abs(w / h - 16 / 9.0) > 0.02:
        raise TemplateError("The template is %.2f x %.2f in, not 16:9. New decks need a widescreen template; build without it, or edit an existing deck in it instead." % (w, h))
    notes = []
    lst = prs.slides._sldIdLst
    count = len(lst)
    for entry in list(lst):  # sample slides in the template are not part of the new deck
        prs.part.drop_rel(entry.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"))
        lst.remove(entry)
    for ext in prs.part._element.xpath(".//*[local-name()='sectionLst']"):
        ext.getparent().getparent().remove(ext.getparent())
    if count:
        notes.append("Removed %d sample slide%s from the template." % (count, "" if count == 1 else "s"))
    colors, fonts = theme_colors_fonts(prs)
    scale = w / DESIGN_W
    layout = blank_layout(prs)
    return prs, theme_tokens(colors, fonts), scale, marks(prs, layout, scale), notes
