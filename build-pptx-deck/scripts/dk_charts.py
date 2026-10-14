"""Native PowerPoint charts and tables, styled from the theme. Both stay editable in PowerPoint."""

import re

from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION,
                             XL_TICK_MARK)
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

import dk_measure as M
from dk_canvas import LH, parse_runs, plain
from dk_theme import contrast, mix, pick_on, readable

TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "stacked_column": XL_CHART_TYPE.COLUMN_STACKED, "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
    "stacked100_column": XL_CHART_TYPE.COLUMN_STACKED_100, "stacked100_bar": XL_CHART_TYPE.BAR_STACKED_100,
    "line": XL_CHART_TYPE.LINE_MARKERS, "area": XL_CHART_TYPE.AREA, "stacked_area": XL_CHART_TYPE.AREA_STACKED,
    "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT, "scatter": XL_CHART_TYPE.XY_SCATTER,
    "waterfall": XL_CHART_TYPE.COLUMN_STACKED,
}
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def _c(hexstr):
    return RGBColor.from_string(hexstr)


def _auto_format(values):
    flat = [v for v in values if isinstance(v, (int, float))]
    if not flat:
        return "General"
    if all(abs(v - round(v)) < 1e-9 for v in flat):
        return "#,##0"
    return "#,##0.0" if max(abs(v) for v in flat) >= 10 else "0.0#"


def _fmt_value(v, fmt):
    """Small formatter for custom labels (waterfall): handles prefix/suffix in quotes, %, decimals."""
    prefix = "".join(re.findall(r'^"([^"]*)"', fmt))
    suffix = "".join(re.findall(r'"([^"]*)"$', fmt))
    core = re.sub(r'"[^"]*"', "", fmt)
    decimals = len(core.split(".")[1].replace("#", "0").rstrip("%")) if "." in core else 0
    sign = "-" if v < 0 else ""
    v = abs(v)
    if "%" in core:
        return "%s%s%.*f%%%s" % (sign, prefix, decimals, v * 100, suffix)
    body = ("{:,.%df}" % decimals).format(v) if "," in core else ("%.*f" % (decimals, v))
    return sign + prefix + body + suffix


def series_colors(c, n):
    out = []
    for i in range(n):
        col = c.P["series"][i % len(c.P["series"])]
        if contrast(col, c.P["bg"]) < 1.35:
            col = readable(col, c.P["bg"], 1.8)
        out.append(col)
    return out


def add_chart(c, x, y, w, h, spec):
    """Add a native chart. See references/layout-catalog.md for the chart spec fields."""
    P = c.P
    kind = spec.get("type", "column")
    if kind not in TYPES:
        c.warn("Unknown chart type %r; used column. Types: %s" % (kind, ", ".join(sorted(TYPES))))
        kind = "column"
    series = spec.get("series") or []
    cats = spec.get("categories") or []
    if not series:
        c.warn("Chart has no series; nothing drawn.")
        return None
    family = c.font("body")
    size = spec.get("font_size", 12)
    labels = spec.get("labels", True)
    all_values = [v for s in series for v in (s.get("values") or []) if isinstance(v, (int, float))]
    fmt = spec.get("number_format") or _auto_format(all_values)

    if kind == "scatter":
        data = XyChartData()
        for s in series:
            ser = data.add_series(s.get("name", ""))
            for pt in s.get("points", []):
                ser.add_data_point(pt[0], pt[1])
    elif kind == "waterfall":
        data, wf = _waterfall_data(cats, series[0].get("values", []), spec.get("totals", []))
    else:
        data = CategoryChartData()
        data.categories = cats
        for s in series:
            vals = list(s.get("values") or [])
            if len(vals) != len(cats):
                c.warn("Chart series %r has %d values for %d categories." % (s.get("name"), len(vals), len(cats)))
            data.add_series(s.get("name", ""), vals)

    frame = c.slide.shapes.add_chart(TYPES[kind], Inches(x), Inches(y), Inches(w), Inches(h), data)
    frame.name = spec.get("name", "Chart")
    chart = frame.chart
    chart.has_title = False
    chart.font.name = family
    chart.font.size = Pt(size)
    chart.font.color.rgb = _c(P["ink"])
    colors = series_colors(c, max(len(series), len(cats)))
    plot = chart.plots[0]
    multi = len(series) > 1 and kind != "waterfall"
    legend = spec.get("legend", "auto")
    show_legend = (multi or kind in ("pie", "doughnut")) if legend == "auto" else bool(legend)
    chart.has_legend = show_legend
    if show_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM if legend in ("auto", True, "bottom") else XL_LEGEND_POSITION.RIGHT
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(size)
        chart.legend.font.color.rgb = _c(P["muted"])
        chart.legend.font.name = family

    if kind in ("pie", "doughnut"):
        plot.vary_by_categories = True
        ser = plot.series[0]
        for i in range(len(cats)):
            pt = ser.points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = _c(colors[i % len(colors)])
            pt.format.line.color.rgb = _c(P["bg"])
            pt.format.line.width = Pt(1.5)
        if labels:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.show_percentage = spec.get("percent", True)
            dl.show_value = not spec.get("percent", True)
            dl.show_category_name = False
            dl.number_format = "0%" if spec.get("percent", True) else fmt
            dl.number_format_is_linked = False
            dl.font.size = Pt(size)
            dl.font.bold = True
            dl.font.color.rgb = _c(P["bg"] if kind == "doughnut" else pick_on(colors[0], "141414", "FFFFFF"))
            if kind == "pie":
                dl.position = XL_LABEL_POSITION.INSIDE_END
        if kind == "doughnut":
            dn = plot._element
            hole = dn.find(qn("c:holeSize"))
            if hole is None:
                hole = dn.makeelement(qn("c:holeSize"), {})
                dn.append(hole)
            hole.set("val", str(spec.get("hole", 62)))
        return frame

    # ---- axes ------------------------------------------------------------------------------
    grid = mix(P["line"], P["bg"], 0.35)
    if kind == "scatter":
        axes = [chart.category_axis, chart.value_axis]
    else:
        axes = [chart.category_axis, chart.value_axis]
    cat_ax, val_ax = axes
    for ax in axes:
        ax.tick_labels.font.size = Pt(size)
        ax.tick_labels.font.color.rgb = _c(P["muted"])
        ax.tick_labels.font.name = family
        ax.major_tick_mark = XL_TICK_MARK.NONE
        ax.minor_tick_mark = XL_TICK_MARK.NONE
    cat_ax.has_major_gridlines = kind == "scatter"
    cat_ax.format.line.color.rgb = _c(P["line"])
    val_ax.has_major_gridlines = True
    val_ax.major_gridlines.format.line.color.rgb = _c(grid)
    val_ax.major_gridlines.format.line.width = Pt(0.75)
    val_ax.format.line.fill.background()
    val_ax.tick_labels.number_format = spec.get("axis_format", fmt)
    val_ax.tick_labels.number_format_is_linked = False
    if kind == "scatter":
        cat_ax.major_gridlines.format.line.color.rgb = _c(grid)
        cat_ax.tick_labels.number_format = spec.get("x_format", "General")
        cat_ax.tick_labels.number_format_is_linked = False
    if "max" in spec:
        val_ax.maximum_scale = spec["max"]
    if "min" in spec:
        val_ax.minimum_scale = spec["min"]
    bars = kind in ("column", "bar", "stacked_column", "stacked_bar", "stacked100_column", "stacked100_bar", "waterfall")
    hide_axis = spec.get("hide_value_axis", labels and bars and len(series) <= 2 and "stacked" not in kind)
    if hide_axis and kind != "scatter":
        # Keep the axis in the file (some renderers misplace negative bars without it) but draw none of it.
        val_ax.tick_label_position = XL_TICK_LABEL_POSITION.NONE
        val_ax.has_major_gridlines = False
    if bars and any(v < 0 for v in all_values):
        cat_ax.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    if kind in ("bar", "stacked_bar", "stacked100_bar"):
        cat_ax.reverse_order = True  # first category on top
        val_ax.tick_label_position  # keep default

    # ---- series ----------------------------------------------------------------------------
    highlight = spec.get("highlight")
    neutral = mix(P["ink"], P["bg"], 0.74)
    if bars:
        for ser in plot.series:
            ser.invert_if_negative = False  # otherwise PowerPoint draws negative bars white
    if kind == "waterfall":
        _style_waterfall(c, plot, wf, fmt, size, spec)
        return frame
    if bars:
        plot.gap_width = spec.get("gap", 55 if not multi else 70)
        plot.overlap = 100 if "stacked" in kind else -8
    for i, ser in enumerate(plot.series):
        col = clean_color(series[i].get("color")) if i < len(series) and series[i].get("color") else colors[i]
        if bars or kind in ("area", "stacked_area"):
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = _c(col)
            ser.format.line.fill.background()
            if bars and not multi and highlight is not None:
                ser.format.fill.fore_color.rgb = _c(neutral)
                for j, cat in enumerate(cats):
                    if cat == highlight or j == highlight or (isinstance(highlight, list) and (cat in highlight or j in highlight)):
                        pt = ser.points[j]
                        pt.format.fill.solid()
                        pt.format.fill.fore_color.rgb = _c(colors[0])
        elif kind == "line":
            ser.smooth = False
            ser.format.line.color.rgb = _c(col)
            ser.format.line.width = Pt(2.5 if (highlight is None or series[i].get("name") == highlight) else 1.5)
            if highlight is not None and series[i].get("name") != highlight:
                ser.format.line.color.rgb = _c(neutral)
                col = neutral
            ser.marker.style = XL_MARKER_STYLE.CIRCLE
            ser.marker.size = 6
            ser.marker.format.fill.solid()
            ser.marker.format.fill.fore_color.rgb = _c(col)
            ser.marker.format.line.color.rgb = _c(P["bg"])
        elif kind == "scatter":
            ser.format.line.fill.background()
            ser.marker.style = XL_MARKER_STYLE.CIRCLE
            ser.marker.size = spec.get("marker", 9)
            ser.marker.format.fill.solid()
            ser.marker.format.fill.fore_color.rgb = _c(col)
            ser.marker.format.line.color.rgb = _c(P["bg"])
    if labels and kind != "scatter":
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = fmt
        dl.number_format_is_linked = False
        dl.font.size = Pt(size)
        dl.font.name = family
        dl.font.color.rgb = _c(P["ink"])
        if "stacked" in kind:
            dl.position = XL_LABEL_POSITION.CENTER
            dl.font.color.rgb = _c(pick_on(colors[0], "141414", "FFFFFF"))
        elif bars:
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
            dl.font.bold = not multi
        elif kind == "line":
            dl.position = XL_LABEL_POSITION.ABOVE
    return frame


def clean_color(c):
    return str(c).strip().lstrip("#").upper()


def _waterfall_data(cats, values, totals):
    """Stacked columns that draw a bridge, including steps that cross or sit below zero.

    Each bar spans lo..hi. Above zero it is an invisible base plus a visible part; below zero the same
    with negative values; a bar that crosses zero is split into a positive and a negative part.
    """
    keys = ("base", "up_pos", "up_neg", "down_pos", "down_neg", "total_pos", "total_neg")
    cols = {k: [] for k in keys}
    labels, kinds = [], []
    level = 0.0
    for i, (cat, v) in enumerate(zip(cats, values)):
        v = float(v or 0)
        if i in totals or cat in totals:
            lo, hi, kind = min(0.0, v), max(0.0, v), "total"
            level = v
        else:
            start, end = level, level + v
            lo, hi, kind = min(start, end), max(start, end), ("up" if v >= 0 else "down")
            level = end
        row = dict.fromkeys(keys, 0.0)
        if lo >= 0:
            row["base"], row[kind + "_pos"] = lo, hi - lo
        elif hi <= 0:
            row["base"], row[kind + "_neg"] = hi, lo - hi
        else:
            row[kind + "_pos"], row[kind + "_neg"] = hi, lo
        for k in keys:
            cols[k].append(row[k])
        labels.append(v)
        kinds.append(kind)
    data = CategoryChartData()
    data.categories = cats
    for k in keys:
        data.add_series(k, cols[k])
    return data, {"labels": labels, "kinds": kinds, "cols": cols, "keys": keys}


def _style_waterfall(c, plot, wf, fmt, size, spec):
    P = c.P
    plot.gap_width = spec.get("gap", 45)
    plot.overlap = 100
    cols = series_colors(c, 3)
    colour = {"up": spec.get("up_color", cols[1] if len(cols) > 1 else cols[0]),
              "down": spec.get("down_color", readable(mix(P["ink"], P["bg"], 0.55), P["bg"], 1.6)),
              "total": spec.get("total_color", cols[0])}
    for idx, key in enumerate(wf["keys"]):
        ser = plot.series[idx]
        if key == "base":
            ser.format.fill.background()
        else:
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = _c(colour[key.split("_")[0]])
        ser.format.line.fill.background()
    for j, kind in enumerate(wf["kinds"]):
        pos, neg = wf["cols"][kind + "_pos"][j], wf["cols"][kind + "_neg"][j]
        key = kind + ("_pos" if abs(pos) >= abs(neg) else "_neg")
        if not (pos or neg):
            continue
        point = plot.series[wf["keys"].index(key)].points[j]
        frame = point.data_label.text_frame
        raw = wf["labels"][j]
        frame.text = ("+" if kind == "up" else "") + _fmt_value(raw, fmt)
        piece = frame.paragraphs[0].runs[0]
        font = piece.font
        font.size = Pt(size)
        font.bold = True
        font.color.rgb = _c(pick_on(colour[kind], "141414", "FFFFFF"))
        point.data_label.position = XL_LABEL_POSITION.CENTER


# ---- tables ----------------------------------------------------------------------------------
def _border(cell, side, color, width):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = {"l": "a:lnL", "r": "a:lnR", "t": "a:lnT", "b": "a:lnB"}[side]
    for el in tcPr.findall(qn(tag)):
        tcPr.remove(el)
    ln = tcPr.makeelement(qn(tag), {"w": str(int(Pt(width))) if color else "0", "cap": "flat", "cmpd": "sng", "algn": "ctr"})
    if color:
        fill = ln.makeelement(qn("a:solidFill"), {})
        fill.append(fill.makeelement(qn("a:srgbClr"), {"val": color}))
        ln.append(fill)
    else:
        ln.append(ln.makeelement(qn("a:noFill"), {}))
    order = ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]
    pos = 0
    for el in list(tcPr):
        if el.tag in [qn(t) for t in order[: order.index(tag)]]:
            pos = list(tcPr).index(el) + 1
    tcPr.insert(pos, ln)


def _looks_numeric(v):
    if plain(str(v)).strip().lower() in ("n/a", "na", "-", "\u2013", "n.a."):
        return True
    return bool(re.match(r"^[\s$\u20ac\u00a3\u20b9+\-\u2212~\u2248<>]*[\d.,]+\s*[%xkKmMbBnT]*\w{0,4}$", plain(str(v)).strip())) if str(v).strip() else False


def add_table(c, x, y, w, spec, max_h):
    """Add a native table sized to its text. Returns the y below it."""
    P = c.P
    cols = [str(v) for v in spec.get("columns", [])]
    rows = [[("" if v is None else str(v)) for v in r] for r in spec.get("rows", [])]
    ncol = max([len(cols)] + [len(r) for r in rows]) if (cols or rows) else 0
    if ncol == 0:
        c.warn("Table has no columns.")
        return y
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    cols = cols + [""] * (ncol - len(cols)) if cols else []
    weights = spec.get("widths") or ([1.7] + [1.0] * (ncol - 1))
    weights = list(weights) + [1.0] * (ncol - len(weights))
    widths = [w * wt / sum(weights[:ncol]) for wt in weights[:ncol]]
    align = spec.get("align") or []
    aligns = []
    for j in range(ncol):
        if j < len(align) and align[j] in ("l", "c", "r"):
            aligns.append(align[j])
        else:
            col_vals = [r[j] for r in rows if r[j]]
            aligns.append("r" if j > 0 and col_vals and all(_looks_numeric(v) for v in col_vals) else "l")
    family = c.font("body")
    padx, pady = 0.12, 0.085
    inverted = spec.get("header", "rule" if c.theme["card"] in ("outline", "rule", "fill") else "fill") == "fill"
    hl_row, hl_col = spec.get("highlight_row"), spec.get("highlight_col")
    size = float(spec.get("size", 14))

    def heights(sz):
        out = []
        if cols:
            hh = max(M.text_height(plain(v).upper(), family, sz - 2, widths[j] - 2 * padx, True, LH, 0, int((sz - 2) * 4))[0]
                     for j, v in enumerate(cols))
            out.append(hh + 2 * pady + 0.02)
        for r in rows:
            rh = max(M.text_height(plain(v), family, sz, widths[j] - 2 * padx, j == 0 and spec.get("first_col_bold", True), LH)[0]
                     for j, v in enumerate(r))
            out.append(rh + 2 * pady)
        return out

    hs = heights(size)
    while sum(hs) > max_h and size > 9:
        size -= 0.5
        hs = heights(size)
    if sum(hs) < max_h * 0.62 and hs:  # a short table gets taller rows instead of floating in space
        grow = min((max_h * 0.62) / sum(hs), 0.62 / max(hs), 1.9)
        if grow > 1:
            hs = [h * grow for h in hs]
    if size < 11:
        c.warn("Table text shrank to %.1f pt to fit: cut rows or columns, shorten cells to a few words, or split the table across slides." % size)
    if sum(hs) > max_h + 0.02:
        c.warn("Table needs %.2f in of height but has %.2f in at 9 pt: cut rows, shorten cells, or split the slide."
               % (sum(hs), max_h))
    total_rows = len(rows) + (1 if cols else 0)
    frame = c.slide.shapes.add_table(total_rows, ncol, Inches(x), Inches(y), Inches(w), Inches(sum(hs)))
    frame.name = spec.get("name", "Table")
    tbl = frame.table
    tblPr = tbl._tbl.tblPr
    for attr in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
        tblPr.set(attr, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = tblPr.makeelement(qn("a:tableStyleId"), {})
        tblPr.append(sid)
    sid.text = NO_STYLE_TABLE
    for j in range(ncol):
        tbl.columns[j].width = Inches(widths[j])
    hl_fill = mix(P["accent"], P["bg"], 0.78)
    for i in range(total_rows):
        tbl.rows[i].height = Inches(hs[i])
        is_head = bool(cols) and i == 0
        data_i = i - (1 if cols else 0)
        for j in range(ncol):
            cell = tbl.cell(i, j)
            value = cols[j] if is_head else rows[data_i][j]
            cell.margin_left = cell.margin_right = Inches(padx)
            cell.margin_top = cell.margin_bottom = Inches(pady)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            fill, ink = None, P["ink"]
            if is_head and inverted:
                fill, ink = P["ink"], P["bg"]
            elif not is_head and (data_i == hl_row or j == hl_col):
                fill = hl_fill
            elif not is_head and spec.get("zebra") and data_i % 2 == 1:
                fill = mix(P["surface"], P["ink"], 0.03) if P["surface"] != P["bg"] else mix(P["bg"], P["ink"], 0.04)
            if fill:
                cell.fill.solid()
                cell.fill.fore_color.rgb = _c(fill)
            else:
                cell.fill.background()
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[aligns[j]]
            for piece, is_bold in parse_runs(value):
                run = p.add_run()
                label = piece.upper() if is_head else piece
                run.text = label
                font = run.font
                font.name = family
                font.size = Pt(size - 2 if is_head else size)
                font.bold = bool(is_head or is_bold or (j == 0 and spec.get("first_col_bold", True)))
                font.color.rgb = _c(ink if (inverted or not is_head) else P["muted"])
                if is_head:
                    run._r.get_or_add_rPr().set("spc", str(int((size - 2) * 4)))
            for side in ("l", "r", "t"):
                _border(cell, side, None, 0)
            last = i == total_rows - 1
            if is_head and not inverted:
                _border(cell, "b", P["ink"], 1.25)
            elif last:
                _border(cell, "b", P["line"] if not spec.get("close", True) else P["ink"], 1.0)
            else:
                _border(cell, "b", P["line"], 0.75)
            c._rec("cell", x + sum(widths[:j]), y + sum(hs[:i]), widths[j], hs[i], size=size, color=ink,
                   on=fill or P["bg"], text=plain(value)[:40], need=0, fits=True, name="cell")
    return y + sum(hs)
